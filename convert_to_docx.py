"""Convert Markdown reports to DOCX format."""
import re
import os
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_cell_shading(cell, color):
    """Set cell background color."""
    shading_elm = OxmlElement("w:shd")
    shading_elm.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shading_elm)


def add_paragraph(doc, text, style_name=None, bold=False, font_size=None, alignment=None, space_after=None):
    """Add a paragraph with optional formatting."""
    p = doc.add_paragraph(style=style_name)
    if alignment is not None:
        p.alignment = alignment
    if space_after is not None:
        pf = p.paragraph_format
        pf.space_after = Pt(space_after)

    run = p.add_run(text)
    if bold:
        run.bold = True
    if font_size:
        run.font.size = Pt(font_size)
    return p


def add_code_block(doc, code_text):
    """Add a styled code block."""
    for line in code_text.strip().split('\n'):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.left_indent = Cm(1)
        run = p.add_run(line)
        run.font.name = 'Consolas'
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)
        # Set background via shading
        shading_elm = OxmlElement("w:shd")
        shading_elm.set(qn("w:fill"), "F0F0F5")
        p._element.get_or_add_pPr().append(shading_elm)


def process_inline_formatting(paragraph, text):
    """Parse inline markdown (**bold**, `code`) and add runs to paragraph."""
    # Pattern: **text** or `text`
    pattern = r'(\*\*(.+?)\*\*|`(.+?)`)'
    last_end = 0
    for match in re.finditer(pattern, text):
        # Add text before match
        before = text[last_end:match.start()]
        if before:
            paragraph.add_run(before)

        if match.group(2):  # **bold**
            run = paragraph.add_run(match.group(2))
            run.bold = True
        elif match.group(3):  # `code`
            run = paragraph.add_run(match.group(3))
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)

        last_end = match.end()

    # Add remaining text
    remaining = text[last_end:]
    if remaining:
        paragraph.add_run(remaining)


def add_rich_paragraph(doc, text, style_name=None, alignment=None, space_after=None):
    """Add a paragraph with inline formatting support."""
    p = doc.add_paragraph(style=style_name)
    if alignment is not None:
        p.alignment = alignment
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    process_inline_formatting(p, text)
    return p


def add_table_from_md(doc, lines, start_idx):
    """Parse a markdown table and add it to the document."""
    # Collect all table rows
    rows = []
    i = start_idx
    while i < len(lines) and '|' in lines[i] and lines[i].strip().startswith('|'):
        row_cells = [c.strip() for c in lines[i].split('|')[1:-1]]
        rows.append(row_cells)
        i += 1

    # Filter separator row
    data_rows = []
    for row in rows:
        if all(re.match(r'^[-:]+$', c) for c in row):
            continue
        data_rows.append(row)

    if not data_rows:
        return i

    num_cols = max(len(r) for r in data_rows)
    table = doc.add_table(rows=len(data_rows), cols=num_cols)
    table.style = 'Light Grid Accent 1'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for ri, row in enumerate(data_rows):
        for ci, cell_text in enumerate(row):
            if ci < num_cols:
                cell = table.cell(ri, ci)
                cell.text = ''
                p = cell.paragraphs[0]
                process_inline_formatting(p, cell_text)
                # Header row styling
                if ri == 0:
                    for run in p.runs:
                        run.bold = True
                        run.font.size = Pt(9)
                else:
                    for run in p.runs:
                        run.font.size = Pt(9)

    doc.add_paragraph()  # spacing after table
    return i


def convert_md_to_docx(md_path, docx_path):
    """Convert a Markdown file to DOCX."""
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    lines = content.split('\n')
    doc = Document()

    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)

    # Configure heading styles
    for level, size in [(1, 22), (2, 16), (3, 13), (4, 11)]:
        heading_style = doc.styles[f'Heading {level}']
        heading_style.font.size = Pt(size)
        heading_style.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)
        if level <= 2:
            heading_style.font.bold = True

    i = 0
    in_code_block = False
    code_lines = []
    in_table = False

    while i < len(lines):
        line = lines[i]

        # Code block (fenced)
        if line.strip().startswith('```'):
            if in_code_block:
                # End code block
                if code_lines:
                    add_code_block(doc, '\n'.join(code_lines))
                    code_lines = []
                    doc.add_paragraph()  # spacing
                in_code_block = False
            else:
                in_code_block = True
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Skip empty lines
        if not line.strip():
            i += 1
            continue

        # Horizontal rule
        if line.strip() in ('---', '***', '___'):
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.space_before = Pt(6)
            pf.space_after = Pt(6)
            # Add a horizontal line via bottom border
            pPr = p._element.get_or_add_pPr()
            pBdr = OxmlElement('w:pBdr')
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), '6')
            bottom.set(qn('w:space'), '1')
            bottom.set(qn('w:color'), 'CCCCCC')
            pBdr.append(bottom)
            pPr.append(pBdr)
            i += 1
            continue

        # Table detection
        if line.strip().startswith('|') and i + 1 < len(lines) and lines[i + 1].strip().startswith('|') and re.match(r'^[\|\s\-:]+$', lines[i + 1].strip()):
            i = add_table_from_md(doc, lines, i)
            continue

        # Heading 1 (# )
        if line.startswith('# ') and not line.startswith('## '):
            text = line[2:].strip()
            add_paragraph(doc, text, style_name='Heading 1')
            i += 1
            continue

        # Heading 2 (## )
        if line.startswith('## ') and not line.startswith('### '):
            text = line[3:].strip()
            add_paragraph(doc, text, style_name='Heading 2')
            i += 1
            continue

        # Heading 3 (### )
        if line.startswith('### ') and not line.startswith('#### '):
            text = line[4:].strip()
            add_paragraph(doc, text, style_name='Heading 3')
            i += 1
            continue

        # Heading 4 (#### )
        if line.startswith('#### '):
            text = line[5:].strip()
            add_paragraph(doc, text, style_name='Heading 4')
            i += 1
            continue

        # Blockquote
        if line.startswith('> '):
            text = line[2:].strip()
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Cm(1)
            run = p.add_run(text)
            run.font.italic = True
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            i += 1
            continue

        # Unordered list
        if re.match(r'^[\-\*]\s+', line):
            text = re.sub(r'^[\-\*]\s+', '', line)
            p = doc.add_paragraph(style='List Bullet')
            p.clear()
            process_inline_formatting(p, text)
            i += 1
            continue

        # Ordered list
        if re.match(r'^\d+\.\s+', line):
            text = re.sub(r'^\d+\.\s+', '', line)
            p = doc.add_paragraph(style='List Number')
            p.clear()
            process_inline_formatting(p, text)
            i += 1
            continue

        # Regular paragraph
        p = doc.add_paragraph()
        process_inline_formatting(p, line.strip())

        # Merge consecutive non-empty lines into same paragraph
        j = i + 1
        while j < len(lines) and lines[j].strip() and not lines[j].startswith('#') and not lines[j].startswith('```') and not lines[j].startswith('|') and not lines[j].startswith('>') and not re.match(r'^[\-\*\d]', lines[j].strip()) and lines[j].strip() not in ('---', '***', '___'):
            p.add_run(' ')
            process_inline_formatting(p, lines[j].strip())
            j += 1
        i = j

    # Set page margins
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    doc.save(docx_path)
    print(f"Saved: {docx_path}")


if __name__ == "__main__":
    base_dir = os.path.dirname(os.path.abspath(__file__))
    reports_dir = os.path.join(base_dir, "reports")

    cn_md = os.path.join(reports_dir, "T2P_Project_Report_CN.md")
    en_md = os.path.join(reports_dir, "T2P_Project_Report_EN.md")

    cn_docx = os.path.join(reports_dir, "T2P_Project_Report_CN.docx")
    en_docx = os.path.join(reports_dir, "T2P_Project_Report_EN.docx")

    for md_path, docx_path in [(cn_md, cn_docx), (en_md, en_docx)]:
        if os.path.exists(md_path):
            convert_md_to_docx(md_path, docx_path)
        else:
            print(f"Not found: {md_path}")

    print("Conversion complete!")
